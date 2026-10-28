"""<AD>_TYBS_metni_tam.md -> <AD>_Arastirma_Onerisi.docx (TÜBİTAK 2209-A Araştırma Önerisi Formu kalıbı).
Çalıştırma: python <beceri>/scripts/docx_uret.py [proje_kökü]   (önce derle.py)
proje.json'dan okunanlar: ad, baslik, yazar, ek2 (EK-2 satırı, isteğe bağlı).
Logo: proje kökünde uretim/tubitak_logo.jpg varsa kapak üstüne konur (logo repoya dahil değildir)."""
import json, os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

KOK = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.getcwd())
AYAR = json.load(open(os.path.join(KOK, 'proje.json'), encoding='utf-8'))
MD = os.path.join(KOK, f"{AYAR['ad']}_TYBS_metni_tam.md")
CIKTI = os.path.join(KOK, f"{AYAR['ad']}_Arastirma_Onerisi.docx")
LOGO = os.path.join(KOK, 'uretim', 'tubitak_logo.jpg')
FONT, BOY = 'Arial', 9
KUTU_GENISLIK = Cm(15.4)

# TÜBİTAK formundaki başlık açıklamaları (2026 formu)
ACIKLAMA = {
    'ÖZET': ['Araştırma önerisi özetinin (1) bilimsel nitelik; (2) yöntem; (3) proje yönetimi ve (4) yaygın etki hakkında bilgileri kapsaması beklenir. Bu bölümün en son yazılması önerilir.'],
    '1.1.': ['Araştırma önerisinde ele alınan konunun kapsamı ve sınırları ile önemi literatürün eleştirel bir değerlendirmesinin yanı sıra nitel veya nicel verilerle açıklanır.',
             'Araştırma önerisinin bilimsel niteliği yazılırken önerinin hangi eksikliği nasıl gidereceği veya hangi soruna nasıl bir çözüm geliştireceği literatüre atıf yapılarak açıklanır. Önerilen çalışmanın araştırma sorusu ve varsa hipotezi açık bir şekilde ortaya konulur.'],
    '1.2.': ['Araştırma önerisinin amacı ve hedefleri açık, ölçülebilir, gerçekçi ve ulaşılabilir nitelikte olacak şekilde yazılır.'],
    '2.': ['Araştırma önerisinde uygulanacak yöntem ve araştırma teknikleri (veri toplama araçları ve analiz yöntemleri dahil) ilgili literatüre atıf yapılarak açıklanır.',
           'Yöntem bölümünün; araştırma tasarımı, bağımlı ve bağımsız değişkenler, istatistiksel yöntemler vb. unsurları içermesi gerekir. Araştırma önerisinde herhangi bir ön çalışma veya fizibilite yapıldıysa bunların sunulması beklenir. Araştırma önerisinde sunulan yöntemlerin çalışma takvimi ile ilişkilendirilmesi gerekir.'],
    '3.1.': ['Araştırmada yer alacak başlıca faaliyetler, her bir faaliyetin hangi sürede gerçekleştirileceği, başarı ölçütü ve araştırmanın başarısına katkısı "Çalışma Takvimi" doldurularak sunulur.'],
    '3.2.': ['Araştırmanın başarısını olumsuz yönde etkileyebilecek riskler ve bu risklerle karşılaşıldığında araştırmanın başarıyla yürütülmesini sağlamak için alınacak tedbirler (B Planı) aşağıdaki Risk Yönetimi Tablosu\'nda ifade edilir. B Plan(lar)ının uygulanması araştırmanın temel hedeflerinden sapmaya yol açmamalıdır.'],
    '3.3.': ['Bu bölümde projenin yürütüleceği kurum ve kuruluşlarda var olan ve projede kullanılacak olan altyapı/ekipman (laboratuvar, araç, makine-teçhizat, vb.) olanakları belirtilir.'],
    '4.': ['Araştırma önerisi kapsamındaki çalışmadan elde edilmesi öngörülen çıktılar amaçlarına göre belirlenen kategorilere ayrılarak belirtilir; ölçülebilir ve gerçekçi hedeflere dayandırılır.'],
    '5.': ['Sadece araştırma önerisinin değerlendirilmesine katkı sağlayabilecek bilgi/veri (grafik, tablo, vb.) eklenebilir.'],
}
TABLO_BASLIK = {'3.1.': 'ÇALIŞMA TAKVİMİ (*)', '3.2.': 'RİSK YÖNETİMİ TABLOSU (*)', '3.3.': 'ARAŞTIRMA OLANAKLARI TABLOSU (*)',
                '4.': 'ARAŞTIRMA ÖNERİSİNDEN BEKLENEN YAYGIN ETKİ TABLOSU'}
TABLO_NOT = {'3.1.': ['(*) Çizelgedeki satırlar ve sütunlar gerektiği kadar genişletilebilir ve çoğaltılabilir.',
                      '(**) Literatür taraması, sonuç raporu hazırlama aşamaları, araştırma sonuçlarının paylaşımı ve malzeme alımı ayrı birer iş adımı olarak gösterilmemelidir.',
                      '(***) Başarı ölçütü, ölçülebilir ve izlenebilir nitelikte olacak şekilde nicel veya nitel ölçütlerle (ifade, sayı, yüzde, vb.) belirtilir. Bu sütundaki değerlerin toplamı 100 olmalıdır.'],
             '3.2.': ['(*) Tablodaki satırlar gerektiği kadar genişletilebilir ve çoğaltılabilir.'],
             '3.3.': ['(*) Tablodaki satırlar gerektiği kadar genişletilebilir ve çoğaltılabilir.']}
BUTCE_NOT = ('NOT: Bütçe talebiniz olması halinde hem bu tablonun hem de TÜBİTAK Yönetim Bilgi Sistemi (TYBS) başvuru ekranında '
             'karşınıza gelecek olan bütçe alanlarının doldurulması gerekmektedir. Tabloda girilen rakamlar ile TYBS ekranındaki rakamlar '
             'arasında farklılık olması halinde TYBS ekranındaki veriler dikkate alınır ve başvuru sonrasında değiştirilemez.')


# ---------- yardımcılar ----------
def run_ekle(p, metin, kalin=False, italik=False, boy=BOY, renk=None, ham=False):
    """**kalın** işaretlerini çözerek paragrafa yazar (ham=True ise çözmeden)."""
    for i, parca in enumerate([metin] if ham else re.split(r'\*\*', metin)):
        if not parca:
            continue
        r = p.add_run(parca)
        r.bold = kalin or (i % 2 == 1)
        r.italic = italik
        r.font.size = Pt(boy)
        r.font.name = FONT
        if renk:
            r.font.color.rgb = renk
    return p


def kaynak_ekle(p, metin, boy):
    """Kaynakçadaki URL'leri tıklanabilir köprü yapar (PDF'te satır kırılan bağlantılar da çalışsın diye)."""
    for i, parca in enumerate(re.split(r'(https?://\S+)', metin)):
        if not parca:
            continue
        if i % 2 == 0:
            run_ekle(p, parca, boy=boy, ham=True)
            continue
        h = OxmlElement('w:hyperlink')
        h.set(qn('r:id'), p.part.relate_to(parca, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
        r = run_ekle(p, parca, boy=boy, ham=True).runs[-1]
        h.append(r._r)
        p._p.append(h)
    return p


def bicim(p, hiza=WD_ALIGN_PARAGRAPH.JUSTIFY, once=0, sonra=4):
    pf = p.paragraph_format
    pf.alignment = hiza
    pf.space_before = Pt(once)
    pf.space_after = Pt(sonra)
    pf.line_spacing = 1.1
    return p


def hucre_kenar(hucre, renk='000000', kalinlik=6):
    tcPr = hucre._tc.get_or_add_tcPr()
    kenar = OxmlElement('w:tcBorders')
    for k in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(kalinlik)); e.set(qn('w:color'), renk)
        kenar.append(e)
    tcPr.append(kenar)


def hucre_dolgu(hucre, renk):
    tcPr = hucre._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), renk)
    tcPr.append(shd)


def hucre_bosluk(tablo, ust=60, sol=100):
    tblPr = tablo._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for k, v in (('top', ust), ('left', sol), ('bottom', ust), ('right', sol)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tblPr.append(mar)


def kutu(belge):
    """Formdaki tek hücreli, çerçeveli cevap kutusu."""
    t = belge.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hucre_bosluk(t, 80, 120)
    c = t.cell(0, 0)
    c.width = Cm(16.2)
    c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    return c


def md_tablo(kap, satirlar, genislikler, toplam=Cm(15.2)):
    basliklar = [h.strip() for h in satirlar[0].strip('|').split('|')]
    veri = [[h.strip() for h in s.strip('|').split('|')] for s in satirlar[2:]]
    t = kap.add_table(rows=1 + len(veri), cols=len(basliklar))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hucre_bosluk(t, 40, 80)
    oran = genislikler or [100 / len(basliklar)] * len(basliklar)
    for j, h in enumerate(basliklar):
        c = t.cell(0, j)
        hucre_dolgu(c, 'D9E2F3')
        p = c.paragraphs[0]
        bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
        run_ekle(p, h, kalin=True, boy=8.5)
    for i, satir in enumerate(veri, 1):
        for j, h in enumerate(satir):
            p = t.cell(i, j).paragraphs[0]
            kisa = len(h) < 14 and j > 0
            bicim(p, WD_ALIGN_PARAGRAPH.CENTER if kisa or j == 0 and len(h) < 6 else WD_ALIGN_PARAGRAPH.LEFT, 0, 0)
            run_ekle(p, h, boy=8.5)
    for j, o in enumerate(oran):
        for i in range(len(veri) + 1):
            t.cell(i, j).width = int(toplam * o / 100)
    return t


def bolum_basligi(belge, metin, seviye):
    p = belge.add_paragraph()
    bicim(p, WD_ALIGN_PARAGRAPH.LEFT, 10 if seviye == 2 else 6, 3)
    p.paragraph_format.keep_with_next = True
    run_ekle(p, metin, kalin=True, boy=10 if seviye == 2 else 9.5)


def aciklama(belge, metinler):
    for m in metinler:
        p = belge.add_paragraph()
        bicim(p, sonra=3)
        p.paragraph_format.keep_with_next = True
        run_ekle(p, m, italik=True, boy=8.5, renk=RGBColor(0x40, 0x40, 0x40))


def icerik_yaz(kap, bloklar):
    """Kutu içine paragraf, madde, şekil ve tablo yazar. kap: hücre veya belge."""
    i = 0
    while i < len(bloklar):
        b = bloklar[i]
        if b.startswith('<!-- cols:'):
            ilk, _, kalan = b.partition('\n')
            gen = [float(x) for x in re.findall(r'[\d.]+', ilk)]
            satirlar = kalan.split('\n') if kalan else []
            i += 1
            while i < len(bloklar) and bloklar[i].startswith('|'):
                satirlar += bloklar[i].split('\n'); i += 1
            md_tablo(kap, satirlar, gen)
            p = kap.add_paragraph(); bicim(p, sonra=2)
            continue
        if b.startswith('!['):
            m = re.match(r'!\[(.*?)\]\((.*?)\)', b)
            alt, yol = m.group(1), os.path.join(KOK, m.group(2))
            p = kap.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 2)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(yol, width=KUTU_GENISLIK)
            p = kap.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
            no, aciklama_m = alt.split('.', 1)
            run_ekle(p, f'**{no}.**{aciklama_m}')
        elif b.startswith('- '):
            for satir in b.split('\n'):
                p = kap.add_paragraph(style='List Bullet'); bicim(p, sonra=1)
                run_ekle(p, satir[2:])
        elif b.startswith('**Tablo'):
            p = kap.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 2, 3)
            p.paragraph_format.keep_with_next = True
            run_ekle(p, b)
        else:
            p = kap.add_paragraph(); bicim(p)
            run_ekle(p, b.replace('\n', ' '))
        i += 1


def bloklara_ayir(metin):
    bloklar = [b.strip() for b in re.split(r'\n\s*\n', metin) if b.strip()]
    return [b for b in bloklar if not b.startswith('>')]


# ---------- belge ----------
def uret():
    md = open(MD, encoding='utf-8').read()
    govde, kaynaklar = md.split('## KAYNAKLAR')
    belge = Document()
    st = belge.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(BOY)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    s = belge.sections[0]
    s.page_width, s.page_height, s.orientation = Twips(11906), Twips(16838), WD_ORIENT.PORTRAIT
    s.top_margin, s.right_margin, s.bottom_margin, s.left_margin = Twips(1440), Twips(1344), Twips(1000), Twips(1239)
    s.header_distance, s.footer_distance = Twips(500), Twips(400)
    hp = s.header.paragraphs[0]; bicim(hp, WD_ALIGN_PARAGRAPH.RIGHT, 0, 0)
    run_ekle(hp, '2209/A ÜNİVERSİTE ÖĞRENCİLERİ ARAŞTIRMA PROJELERİ DESTEĞİ PROGRAMI\nARAŞTIRMA ÖNERİSİ FORMU', boy=7, renk=RGBColor(0x55, 0x55, 0x55))
    fp = s.footer.paragraphs[0]; bicim(fp, WD_ALIGN_PARAGRAPH.CENTER, 0, 0)
    r = fp.add_run(); r.font.size = Pt(8)
    for tip, deger in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tip:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tip); r._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = deger; r._r.append(e)

    # Kapak üstü
    p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
    if os.path.exists(LOGO): p.add_run().add_picture(LOGO, width=Cm(2.6))
    p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
    run_ekle(p, 'TÜBİTAK–2209-A ÜNİVERSİTE ÖĞRENCİLERİ ARAŞTIRMA PROJELERİ DESTEĞİ PROGRAMI', kalin=True, boy=11)
    p = belge.add_paragraph(); bicim(p, sonra=8)
    run_ekle(p, 'Başvuru formunun Arial 9 yazı tipinde, her bir konu başlığı altında verilen açıklamalar göz önünde bulundurularak '
                'hazırlanması ve ekler hariç toplam 20 sayfayı geçmemesi beklenir (Alt sınır bulunmamaktadır). Değerlendirme araştırma '
                'önerisinin bilimsel niteliği, yöntemi, proje yönetimi ve yaygın etkisi başlıkları üzerinden yapılacaktır.', kalin=True)
    for m, b in (('ARAŞTIRMA ÖNERİSİ FORMU', 12), ('2026 Yılı', 10)):
        p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 4)
        run_ekle(p, m, kalin=True, boy=b)

    # Bölümleri ayır
    parcalar = re.split(r'^(#{2,3} .+)$', govde, flags=re.M)
    bolumler = [(parcalar[k].lstrip('#').strip(), parcalar[k + 1]) for k in range(1, len(parcalar), 2)]
    for baslik, metin in bolumler:
        anahtar = baslik.split()[0] if baslik[0].isdigit() else baslik
        bloklar = bloklara_ayir(metin)
        if baslik.startswith('A. GENEL'):
            bolum_basligi(belge, baslik, 2)
            satirlar = [s for s in metin.split('\n') if s.startswith('|')][2:]
            t = belge.add_table(rows=0, cols=1); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hucre_bosluk(t, 50, 120)
            for s_ in satirlar[:4]:
                a, d = [x.strip() for x in s_.strip('|').split('|')]
                p = t.add_row().cells[0].paragraphs[0]; bicim(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 0)
                run_ekle(p, f'**{a}:** {d}')
            continue
        bolum_basligi(belge, baslik, 2 if re.match(r'^(\d\. |ÖZET)', baslik) else 3)
        if anahtar in ACIKLAMA:
            aciklama(belge, ACIKLAMA[anahtar])
        if anahtar == 'ÖZET':
            c = kutu(belge)
            p = c.add_paragraph(); bicim(p, sonra=2); run_ekle(p, 'Özet', kalin=True)
            icerik_yaz(c, [b for b in bloklar if not b.startswith('**Anahtar')])
            p = c.add_paragraph(); bicim(p, 4, 0)
            run_ekle(p, [b for b in bloklar if b.startswith('**Anahtar')][0])
        elif anahtar in TABLO_BASLIK:
            p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.CENTER, 2, 3)
            run_ekle(p, TABLO_BASLIK[anahtar], kalin=True)
            icerik_yaz(belge, bloklar)
            for n in TABLO_NOT.get(anahtar, []):
                p = belge.add_paragraph(); bicim(p, sonra=1); run_ekle(p, n, boy=7.5, ham=True)
        elif anahtar == '5.':
            butce_i = next(k for k, b in enumerate(bloklar) if b.startswith('**Bütçe'))
            c = kutu(belge)
            icerik_yaz(c, bloklar[:butce_i])
            bolum_basligi(belge, 'BÜTÇE TALEP ÇİZELGESİ', 3)
            icerik_yaz(belge, bloklar[butce_i + 1:])
            p = belge.add_paragraph(); bicim(p, sonra=2); run_ekle(p, BUTCE_NOT, boy=7.5)
        elif anahtar == '6.':
            p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.LEFT, 2, 4); run_ekle(p, 'EK-1: KAYNAKLAR', kalin=True)
            for k in [k.strip() for k in kaynaklar.split('\n\n') if k.strip()]:
                p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 3)
                p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.7)
                kaynak_ekle(p, k, boy=8.5)
            p = belge.add_paragraph(); bicim(p, WD_ALIGN_PARAGRAPH.LEFT, 6, 0)
            if AYAR.get('ek2'): run_ekle(p, 'EK-2: ' + AYAR['ek2'])
        elif bloklar:
            c = kutu(belge)
            icerik_yaz(c, bloklar)

    belge.core_properties.title = AYAR.get('baslik', 'TÜBİTAK 2209-A Araştırma Önerisi Formu')
    belge.core_properties.author = AYAR.get('yazar', '')
    belge.save(CIKTI)
    print('yazıldı:', CIKTI)


if __name__ == '__main__':
    uret()
